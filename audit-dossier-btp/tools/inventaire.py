#!/usr/bin/env python3
"""Phase 2 — inventaire documentaire et extraction de texte (lecture seule).

Usage :
  python inventaire.py --source <dossier> --travail <_travail>

Produit :
  _travail/manifeste.json            inventaire complet (empreintes, formats, pages, doublons,
                                     familles de versions, références citées, normes citées,
                                     consignes suspectes adressées à une IA)
  _travail/extraits/Dxxx.txt         texte extrait, balisé « === PAGE n (affichée : X) === »
Stdout : résumé compact JSON pour l'orchestrateur.
"""
from __future__ import annotations

import argparse
import datetime
import hashlib
import re
import sys
import zipfile
from collections import defaultdict
from pathlib import Path
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parent))
from btp_common import (journal, norm_text, now_iso, print_summary, read_table,  # noqa: E402
                        save_json)

FORMATS = {
    "pdf": {".pdf"},
    "tableur": {".xlsx", ".xlsm", ".csv", ".tsv"},
    "tableur_ancien": {".xls", ".ods"},
    "traitement_texte": {".docx"},
    "traitement_texte_ancien": {".doc", ".odt", ".rtf"},
    "texte": {".txt", ".md"},
    "maquette_ifc": {".ifc"},
    "maquette_proprietaire": {".rvt", ".rfa", ".nwd", ".nwc", ".pln", ".skp"},
    "dao": {".dwg", ".dxf", ".dwf"},
    "image": {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp"},
    "archive": {".zip", ".rar", ".7z"},
}
EXCLUDED_DIRS = {".git", "_travail", "__pycache__", ".claude"}

PAGE_NUM_PATTERNS = [
    re.compile(r"^\s*(?:page|p\.)\s*(\d{1,4})\s*(?:/|sur)\s*\d{1,4}\s*$", re.I),
    re.compile(r"^\s*(\d{1,4})\s*/\s*\d{1,4}\s*$"),
    re.compile(r"^\s*(?:page|p\.)\s*(\d{1,4})\s*$", re.I),
    re.compile(r"^\s*-\s*(\d{1,4})\s*-\s*$"),
]

REF_PATTERNS = {
    "plan": re.compile(r"\bplans?\s*(?:n°|no|num[ée]ro)?\s*([A-Z]{1,4}[-_.]?\d{2,4}[A-Z]?)\b"),
    "detail": re.compile(r"\bd[ée]tails?\s*(?:n°|no)?\s*([A-Z]{0,3}[-]?\d{1,3}[A-Z]?)\b", re.I),
    "annexe": re.compile(r"\bannexe\s*(?:n°\s*)?([A-Z0-9]{1,4})\b", re.I),
}
NORM_PATTERNS = [
    re.compile(r"\b(?:NF\s+)?DTU\s*\d{1,2}(?:\.\d{1,2})?(?:\s*P\s?\d(?:-\d)?)?", re.I),
    re.compile(r"\bNF\s+EN\s+(?:ISO\s+)?\d{2,5}(?:-\d+)*(?:\+A\d)?", re.I),
    re.compile(r"\bNF\s+[A-Z]\s?\d{2}-\d{3}(?:-\d)?", re.I),
    re.compile(r"\bEurocode\s*\d{1,2}\b", re.I),
    re.compile(r"\bRE\s?2020\b", re.I),
    re.compile(r"\bRT\s?2012\b", re.I),
    re.compile(r"\bCPT\s*\d{3,4}\b", re.I),
    re.compile(r"\barr[êe]t[ée]\s+du\s+\d{1,2}\s+[a-zéû]+\s+\d{4}", re.I),
]
INJECTION_PATTERNS = [
    re.compile(r"ignor(?:e|ez|er)\s+(?:toutes\s+)?(?:les\s+)?(?:pr[ée]c[ée]dentes\s+)?(?:instructions|consignes)", re.I),
    re.compile(r"ignore\s+(?:all\s+)?(?:previous|prior|above)\s+instructions", re.I),
    re.compile(r"\b(?:tu\s+es|vous\s+[êe]tes)\s+(?:un|une)\s+(?:ia|intelligence artificielle|assistant|mod[èe]le)", re.I),
    re.compile(r"\byou\s+are\s+(?:an?\s+)?(?:ai|assistant|language model|chatgpt|claude)\b", re.I),
    re.compile(r"\bsystem\s+prompt\b", re.I),
    re.compile(r"\b(?:ne\s+signale[zr]?\s+pas|do\s+not\s+(?:report|mention|flag))\b", re.I),
    re.compile(r"<\s*/?\s*(?:instructions?|system)\s*>", re.I),
    re.compile(r"\b(?:note|message)\s+(?:pour|à|to)\s+(?:l'|the\s+)?(?:ia|ai|assistant|claude|chatgpt)\b", re.I),
]
LOT_PATTERN = re.compile(r"\bLOT\s*(?:N°\s*|n°\s*)?(\d{1,2}[A-Z]?)\s*[-–:.]?\s*([A-ZÉÈÀÂÎÔÛÇ][A-Za-zÀ-ÿ'’ /&,-]{2,60})")
VERSION_TOKEN = re.compile(r"(?:^|[_\-\s.])(?:ind(?:ice)?|rev|r[ée]v|v|version)[_\-\s.]?([A-Z0-9]{1,3})(?=$|[_\-\s.])", re.I)
INDICE_TEXT = re.compile(r"\bindice\s*[:=]?\s*([A-Z0-9]{1,2})\b", re.I)


def fmt_of(path: Path) -> str:
    ext = path.suffix.lower()
    for name, exts in FORMATS.items():
        if ext in exts:
            return name
    return "autre"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def displayed_page(text: str):
    lines = [ln for ln in text.splitlines() if ln.strip()]
    for ln in lines[-4:] + lines[:3]:
        for pat in PAGE_NUM_PATTERNS:
            m = pat.match(ln)
            if m:
                return m.group(1)
    return None


def extract_pdf(path: Path):
    pages = []
    try:
        import pdfplumber  # type: ignore
        with pdfplumber.open(str(path)) as pdf:
            for pg in pdf.pages[:2000]:
                pages.append(pg.extract_text() or "")
        return pages, "pdfplumber"
    except ImportError:
        pass
    try:
        from pypdf import PdfReader  # type: ignore
        reader = PdfReader(str(path))
        for pg in reader.pages[:2000]:
            pages.append(pg.extract_text() or "")
        return pages, "pypdf"
    except ImportError:
        return None, None


def pdf_page_count_raw(path: Path):
    try:
        data = path.read_bytes()
    except OSError:
        return None
    n = len(re.findall(rb"/Type\s*/Page[^s]", data))
    return n or None


def extract_docx(path: Path):
    try:
        import docx  # type: ignore
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts), "python-docx"
    except ImportError:
        pass
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(path) as z:
        root = ET.fromstring(z.read("word/document.xml"))
    out = []
    for p in root.iter(ns + "p"):
        out.append("".join(t.text or "" for t in p.iter(ns + "t")))
    return "\n".join(out), "stdlib"


def render_table(path: Path):
    sheets, missing = read_table(path)
    blocks = []
    for i, sh in enumerate(sheets, start=1):
        lines = [f"L{r}: " + " | ".join(str(c).strip() for c in cells) for r, cells in sh["rows"]]
        blocks.append((f"feuille {sh['name']}", "\n".join(lines)))
    return blocks, missing


def write_extract(path: Path, pages):
    """pages : liste de (page_affichee|None, texte)."""
    with path.open("w", encoding="utf-8") as fh:
        for i, (disp, text) in enumerate(pages, start=1):
            label = f" (affichée : {disp})" if disp else ""
            fh.write(f"=== PAGE {i}{label} ===\n{text.rstrip()}\n")


def family_key(stem: str):
    m = VERSION_TOKEN.search(stem)
    token = m.group(1).upper() if m else None
    base = VERSION_TOKEN.sub("", stem) if m else stem
    base = re.sub(r"[_\-\s.]+", " ", base).strip().lower()
    return base, token


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", required=True)
    ap.add_argument("--travail", required=True)
    a = ap.parse_args()
    src = Path(a.source).expanduser().resolve()
    work = Path(a.travail).expanduser().resolve()
    out_dir = work.parent
    (work / "extraits").mkdir(parents=True, exist_ok=True)

    if src.is_file():
        files = [src]
        root = src.parent
    else:
        root = src
        files = []
        # la sortie n'est exclue que si elle se trouve DANS le dossier source
        try:
            out_dir.relative_to(root)
            sortie_dans_source = out_dir != root
        except ValueError:
            sortie_dans_source = False
        for p in sorted(src.rglob("*"), key=lambda x: str(x.relative_to(src)).lower()):
            if not p.is_file():
                continue
            rel_parts = p.relative_to(src).parts
            if any(part in EXCLUDED_DIRS or part.startswith("livrables_audit") for part in rel_parts[:-1]):
                continue
            if sortie_dans_source:
                try:
                    p.relative_to(out_dir)
                    continue
                except ValueError:
                    pass
            if p.name.startswith("~$") or p.name in (".DS_Store", "Thumbs.db"):
                continue
            files.append(p)

    docs = []
    by_hash = defaultdict(list)
    norm_refs = defaultdict(set)
    ref_hits = []
    injections = []
    lots = defaultdict(set)
    texts = {}

    for i, p in enumerate(files, start=1):
        doc_id = f"D{i:03d}"
        rel = str(p.relative_to(root)).replace("\\", "/")
        fmt = fmt_of(p)
        st = p.stat()
        rec = {"doc_id": doc_id, "chemin_relatif": rel, "nom": p.name, "extension": p.suffix.lower(),
               "format": fmt, "taille_octets": st.st_size, "sha256": sha256(p),
               "date_modification_fichier": datetime.datetime.fromtimestamp(st.st_mtime).astimezone().isoformat(timespec="seconds"),
               "nb_pages": None, "extraction": None, "moteur_extraction": None,
               "extrait": None, "pages_sans_texte": [], "remarques": []}
        by_hash[rec["sha256"]].append(doc_id)
        pages = None
        try:
            if fmt == "pdf":
                raw_pages, engine = extract_pdf(p)
                if raw_pages is None:
                    rec["nb_pages"] = pdf_page_count_raw(p)
                    rec["extraction"] = "bibliotheque_absente"
                    rec["remarques"].append("texte non extrait (aucune bibliothèque PDF) : lecture directe du PDF requise")
                else:
                    rec["nb_pages"] = len(raw_pages)
                    rec["moteur_extraction"] = engine
                    empty = [k + 1 for k, t in enumerate(raw_pages) if len(t.strip()) < 20]
                    rec["pages_sans_texte"] = empty
                    if raw_pages and len(empty) == len(raw_pages):
                        rec["extraction"] = "scan_sans_texte"
                        rec["remarques"].append("aucune couche texte : document scanné ou dessin vectoriel sans texte, lecture visuelle requise")
                    elif empty:
                        rec["extraction"] = "texte_partiel"
                    else:
                        rec["extraction"] = "texte_extrait"
                    pages = [(displayed_page(t), t) for t in raw_pages]
            elif fmt == "tableur":
                blocks, missing = render_table(p)
                rec["nb_pages"] = len(blocks)
                rec["extraction"] = "tableur_extrait"
                rec["moteur_extraction"] = "openpyxl/stdlib"
                if missing:
                    rec["remarques"].append(f"{missing} cellule(s) contiennent une formule sans valeur en cache : "
                                            "ouvrir et enregistrer le fichier dans Excel pour fiabiliser la lecture")
                pages = blocks
            elif fmt == "traitement_texte":
                text, engine = extract_docx(p)
                rec["nb_pages"] = None
                rec["extraction"] = "texte_extrait"
                rec["moteur_extraction"] = engine
                rec["remarques"].append("document non paginé à l'extraction : localiser par article/paragraphe")
                pages = [(None, text)]
            elif fmt == "texte":
                text = p.read_text(encoding="utf-8", errors="replace")
                chunks = text.split("\f") if "\f" in text else [text]
                rec["nb_pages"] = len(chunks)
                rec["extraction"] = "texte_extrait"
                rec["moteur_extraction"] = "stdlib"
                pages = [(displayed_page(c), c) for c in chunks]
            elif fmt == "maquette_ifc":
                rec["extraction"] = "maquette_ifc"
                rec["remarques"].append("analyser avec ifc_probe.py")
            elif fmt == "image":
                rec["extraction"] = "lecture_visuelle_requise"
            elif fmt == "archive":
                rec["extraction"] = "archive_non_extraite"
                rec["remarques"].append("archive non décompressée par sécurité : la décompresser manuellement puis relancer")
            elif fmt in ("maquette_proprietaire", "dao", "tableur_ancien", "traitement_texte_ancien"):
                rec["extraction"] = "non_pris_en_charge"
                rec["remarques"].append("format non lisible par les outils disponibles : fournir un export PDF, XLSX ou IFC")
            else:
                rec["extraction"] = "non_pris_en_charge"
        except Exception as exc:  # lecture impossible : on le dit, on continue
            rec["extraction"] = "erreur_lecture"
            rec["remarques"].append(f"erreur de lecture : {type(exc).__name__}: {exc}")
            pages = None

        if pages is not None:
            ext_path = work / "extraits" / f"{doc_id}.txt"
            write_extract(ext_path, pages)
            rec["extrait"] = str(ext_path)
            full = "\n".join(t for _, t in pages)
            texts[doc_id] = full
            for n, (_, t) in enumerate(pages, start=1):
                spans = []
                for pat in INJECTION_PATTERNS:
                    for m in pat.finditer(t):
                        if any(abs(m.start() - x) < 300 for x in spans):
                            continue
                        spans.append(m.start())
                        s = max(0, m.start() - 80)
                        injections.append({"doc_id": doc_id, "page_pdf": n,
                                           "extrait": t[s:m.end() + 120].replace("\n", " ")})
                for kind, pat in REF_PATTERNS.items():
                    for m in pat.finditer(t):
                        ref_hits.append({"type": kind, "reference": m.group(1), "cite_dans": doc_id, "page_pdf": n})
                for pat in NORM_PATTERNS:
                    for m in pat.finditer(t):
                        norm_refs[re.sub(r"\s+", " ", m.group(0).strip())].add(doc_id)
                for m in LOT_PATTERN.finditer(t):
                    lots[m.group(1).lstrip("0") or "0"].add(re.sub(r"\s+", " ", m.group(2)).strip(" -–:.,")[:60])
            if rec["extraction"] in ("texte_extrait", "texte_partiel"):
                m = INDICE_TEXT.search(full[:4000])
                if m:
                    rec["indice_detecte_texte"] = m.group(1).upper()
        fam, token = family_key(p.stem)
        rec["famille"] = fam
        rec["indice_detecte_nom"] = token
        docs.append(rec)

    # doublons exacts
    doublons = [ids for ids in by_hash.values() if len(ids) > 1]
    # familles de versions
    fams = defaultdict(list)
    for d in docs:
        fams[(d["famille"], d["format"])].append(d)
    familles = []
    for (fam, fmt), members in fams.items():
        toks = {m["indice_detecte_nom"] or m.get("indice_detecte_texte") for m in members}
        if len(members) > 1:
            familles.append({"famille": fam, "format": fmt,
                             "documents": [{"doc_id": m["doc_id"], "nom": m["nom"],
                                            "indice_nom": m["indice_detecte_nom"],
                                            "indice_texte": m.get("indice_detecte_texte")} for m in members],
                             "indices_distincts": sorted(t for t in toks if t),
                             "commentaire": "versions concurrentes possibles : la pièce applicable n'est PAS déduite automatiquement"})
    # références citées non retrouvées dans les noms/textes d'autres pièces
    names_norm = {d["doc_id"]: norm_text(d["nom"]) for d in docs}
    refs_out = []
    seen = set()
    for h in ref_hits:
        key = (h["type"], h["reference"].upper())
        if key in seen:
            continue
        seen.add(key)
        rnorm = norm_text(h["reference"])
        in_names = [did for did, n in names_norm.items() if rnorm and rnorm in n]
        cites = sorted({x["cite_dans"] for x in ref_hits if (x["type"], x["reference"].upper()) == key})
        in_texts = [did for did, t in texts.items() if did not in cites and rnorm
                    and re.search(r"\b" + re.escape(h["reference"]) + r"\b", t)]
        refs_out.append({**h, "cite_dans": cites, "trouvee_dans_noms": in_names,
                         "presente_dans_autres_textes": in_texts[:10],
                         "statut": "localisee" if in_names else ("mentionnee_ailleurs" if in_texts else "non_retrouvee")})

    manifest = {"genere_le": now_iso(), "source": str(src), "nb_documents": len(docs), "documents": docs,
                "doublons_exacts": doublons, "familles_versions": familles,
                "references_citees": refs_out,
                "normes_citees": [{"reference": k, "documents": sorted(v)} for k, v in sorted(norm_refs.items())],
                "lots_mentionnes": [{"lot": k, "intitules": sorted(v)[:5]} for k, v in
                                    sorted(lots.items(), key=lambda kv: (len(kv[0]), kv[0]))],
                "consignes_suspectes": injections[:50]}
    save_json(work / "manifeste.json", manifest)

    stats = defaultdict(int)
    for d in docs:
        stats[d["extraction"]] += 1
    summary = {"statut": "ok", "outil": "inventaire", "nb_documents": len(docs),
               "par_format": {f: sum(1 for d in docs if d["format"] == f) for f in sorted({d["format"] for d in docs})},
               "par_extraction": dict(stats),
               "doublons_exacts": len(doublons), "familles_versions": len(familles),
               "references_non_retrouvees": sum(1 for r in refs_out if r["statut"] == "non_retrouvee"),
               "normes_citees": len(norm_refs), "consignes_suspectes": len(injections),
               "presence": {"ifc": any(d["format"] == "maquette_ifc" for d in docs),
                            "maquette_proprietaire": any(d["format"] == "maquette_proprietaire" for d in docs),
                            "tableurs": sum(1 for d in docs if d["format"] == "tableur")},
               "documents": [{"doc_id": d["doc_id"], "nom": d["chemin_relatif"], "format": d["format"],
                              "pages": d["nb_pages"], "extraction": d["extraction"]} for d in docs[:150]],
               "manifeste": str(work / "manifeste.json")}
    if len(docs) > 150:
        summary["documents_tronques"] = len(docs) - 150
    journal(work, "inventaire", "ok", {k: summary[k] for k in ("nb_documents", "par_extraction",
                                                                "doublons_exacts", "familles_versions",
                                                                "consignes_suspectes")})
    print_summary(summary)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
