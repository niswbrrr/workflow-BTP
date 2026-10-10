#!/usr/bin/env python3
"""Phase 10 — génération des livrables à partir du registre consolidé (et de lui seul).

Usage :
  python livrables.py --travail <_travail>

Écrit dans le répertoire parent de _travail :
  rapport_audit_dossier_btp.md, registre_anomalies.csv (+ .xlsx si openpyxl),
  rapprochement_economique.csv, questions_a_arbitrer.md, risques_et_interfaces.md,
  manifeste_documentaire.csv, journal_controles.md
CSV : séparateur « ; », encodage UTF-8 avec BOM, virgule décimale (ouverture directe dans Excel FR).
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from btp_common import fmt_fr, journal, load_json, print_summary  # noqa: E402

STATUS_FR = {"corrobore": "Corroboré", "a_verifier": "À vérifier", "conteste": "Contesté", "refute": "Réfuté",
             "non_resolu": "Non résolu"}
TYPE_FR = {"anomalie_averee": "Anomalie avérée", "incoherence": "Incohérence", "omission_possible": "Omission possible",
           "risque": "Risque", "hypothese": "Hypothèse", "question_ouverte": "Question ouverte"}
SEV_FR = {"critique": "Critique", "majeure": "Majeure", "moderee": "Modérée", "mineure": "Mineure"}
CONF_FR = {"elevee": "Élevée", "moyenne": "Moyenne", "faible": "Faible"}
LEVEL_FR = {"prouve_par_outil": "prouvé par contrôle déterministe", "citation_verifiee": "citation vérifiée",
            "observation_localisee": "observation localisée (citation non vérifiable)", "non_etaye": "non étayé"}
SEV_ORDER = {"critique": 0, "majeure": 1, "moderee": 2, "mineure": 3}
CSV_COLS = ["finding_id", "title", "discipline", "lot", "finding_type", "finding_type_initial", "severity",
            "severity_criteria", "severity_rationale", "priority", "decision_deadline", "status", "motif_statut",
            "confidence", "confidence_rationale", "evidence_level", "claim", "evidence", "sources", "calculation",
            "calcul_verification", "economic_impact", "technical_impact", "dependencies", "counterarguments",
            "recommended_action", "responsible_party", "open_questions", "origin_agents", "merged_from",
            "echanges", "verifications", "arbitrage", "quality_flags", "phase_origine"]


ROLES = [  # premier motif trouvé dans l'acteur cité -> rôle canonique
    (("structure",), "Bureau d'études structure"),
    (("fluide", "cvc", "plomberie", "electric", "électric"), "Bureaux d'études et lots techniques"),
    (("econom", "économ", "metreur", "métreur"), "Économiste / métreur"),
    (("cctp", "redacteur", "rédacteur"), "Rédacteur des pièces écrites"),
    (("opc", "planning", "ordonnancement"), "OPC"),
    (("maitre d'ouvrage", "maître d'ouvrage", "maitrise d'ouvrage", "maîtrise d'ouvrage", "moa"), "Maître d'ouvrage"),
    (("entreprise", "titulaire"), "Entreprises"),
    (("architecte", "moe", "maitrise d'oeuvre", "maîtrise d'œuvre", "maîtrise d'oeuvre", "maitre d'oeuvre", "maître d'oeuvre"),
     "Architecte / maîtrise d'œuvre"),
]


def role(acteur):
    a = (acteur or "").lower()
    for motifs, nom in ROLES:
        if any(m in a for m in motifs):
            return nom
    return acteur or "Maîtrise d'œuvre (à préciser)"


def eur(n):
    """Montant pour lecture humaine : « 1 045,00 € » (espace fine insécable pour les milliers)."""
    if n is None:
        return "non chiffré"
    txt = f"{float(n):,.2f}".replace(",", "\u202f").replace(".", ",")
    return f"{txt} €"


def src_str(s):
    parts = [s.get("doc_id", "?")]
    if s.get("page"):
        parts.append(f"p.{s['page']}")
    elif s.get("page_pdf"):
        parts.append(f"p.PDF {s['page_pdf']}")
    if s.get("localisation"):
        parts.append(s["localisation"])
    if s.get("url"):
        parts.append(s["url"])
    txt = " ".join(parts)
    if s.get("citation"):
        txt += f" « {s['citation'][:160]} »"
    if s.get("verification"):
        txt += f" [{s['verification']}]"
    return txt


def calc_str(c):
    if not c:
        return ""
    vs = []
    for k, v in (c.get("variables") or {}).items():
        if isinstance(v, dict):
            vs.append(f"{k}={v.get('valeur')} {v.get('unite') or ''}".strip() + (f" ({v['source']})" if v.get("source") else ""))
        else:
            vs.append(f"{k}={v}")
    s = f"{c.get('formule', '')} ; {', '.join(vs)} => {c.get('resultat')} {c.get('unite') or ''}".strip()
    if c.get("controle_id"):
        s += f" [contrôle {c['controle_id']}]"
    return s


def eco_str(e):
    if not e:
        return "non évalué"
    if e.get("quantifie") and e.get("montant_ht") is not None:
        return f"{eur(e['montant_ht'])} HT ({e.get('base') or 'base non précisée'})"
    return "non quantifié" + (f" : {e['commentaire']}" if e.get("commentaire") else "")


def row_of(f):
    return {
        "finding_id": f["finding_id"], "title": f["title"], "discipline": f.get("discipline"), "lot": f.get("lot"),
        "finding_type": TYPE_FR.get(f["finding_type"], f["finding_type"]),
        "finding_type_initial": TYPE_FR.get(f.get("finding_type_initial"), f.get("finding_type_initial")),
        "severity": SEV_FR.get(f["severity"], f["severity"]),
        "severity_criteria": ", ".join(f.get("severity_criteria") or []),
        "severity_rationale": f.get("severity_rationale"), "priority": f.get("priority"),
        "decision_deadline": f.get("decision_deadline") or "",
        "status": STATUS_FR.get(f["status"], f["status"]), "motif_statut": f.get("motif_statut"),
        "confidence": CONF_FR.get(f["confidence"], f["confidence"]),
        "confidence_rationale": f.get("confidence_rationale"),
        "evidence_level": LEVEL_FR.get(f.get("evidence_level"), f.get("evidence_level")),
        "claim": f["claim"],
        "evidence": " | ".join(f"[{e['nature']}] {e['fait']}" for e in f.get("evidence", [])),
        "sources": " | ".join(src_str(s) for s in f.get("sources", [])),
        "calculation": calc_str(f.get("calculation")), "calcul_verification": f.get("calcul_verification") or "",
        "economic_impact": eco_str(f.get("economic_impact")), "technical_impact": f.get("technical_impact") or "",
        "dependencies": ", ".join(f.get("dependencies", [])),
        "counterarguments": " | ".join(f.get("counterarguments", []) or []),
        "recommended_action": f.get("recommended_action"), "responsible_party": f.get("responsible_party") or "",
        "open_questions": " | ".join(f.get("open_questions", []) or []),
        "origin_agents": ", ".join(f.get("origin_agents", [])), "merged_from": ", ".join(f.get("merged_from", [])),
        "echanges": " | ".join(f"tour {e['round']} question de {e.get('de')}, réponse de {e['a']} : {e['verdict']} — {e['justification'][:200]}"
                               for e in f.get("echanges", [])),
        "verifications": " | ".join(f"{v['lentille']} : {v['verdict']} — {v['justification'][:200]}"
                                    for v in f.get("verifications", [])),
        "arbitrage": "" if not f.get("arbitrage") else f"{f['arbitrage']['decision']} — {f['arbitrage']['motif'][:300]}",
        "quality_flags": ", ".join(f.get("quality_flags", [])), "phase_origine": f.get("phase_origine") or ""}


def write_csv(path, cols, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writerow(cols)
        for r in rows:
            w.writerow([("" if r.get(c) is None else r.get(c)) for c in cols])


def md_escape(s):
    return str(s or "").replace("|", "\\|").replace("\n", " ")


def fid_label(fid, index):
    f = index.get(fid)
    if not f:
        return f"{fid} (identifiant inconnu)"
    return f"{fid} ({STATUS_FR[f['status']].lower()})"


def finding_block(f, index):
    lines = [f"#### {f['finding_id']} — {f['title']}",
             f"- **Type** : {TYPE_FR.get(f['finding_type'])}"
             + (f" (initialement {TYPE_FR.get(f['finding_type_initial'])})" if f['finding_type'] != f.get('finding_type_initial') else "")
             + f" · **Gravité** : {SEV_FR[f['severity']]} · **Priorité** : {f['priority']}"
             + f" · **Statut** : {STATUS_FR[f['status']]} · **Confiance** : {CONF_FR[f['confidence']]}",
             f"- **Lot / discipline** : {f.get('lot') or 'non précisé'} / {f.get('discipline')}",
             f"- **Constat** : {f['claim']}",
             f"- **Preuve** ({LEVEL_FR.get(f.get('evidence_level'))}) : "
             + ("; ".join(src_str(s) for s in f.get("sources", [])[:4]) or "aucune source")]
    if f.get("calculation"):
        lines.append(f"- **Calcul** : {calc_str(f['calculation'])} — vérification : {f.get('calcul_verification')}")
    lines.append(f"- **Impact économique** : {eco_str(f.get('economic_impact'))}")
    if f.get("technical_impact"):
        lines.append(f"- **Impact technique** : {f['technical_impact']}")
    lines.append(f"- **Justification du statut** : {f.get('motif_statut')}")
    if f.get("counterarguments"):
        lines.append(f"- **Contre-arguments** : {' ; '.join(f['counterarguments'])}")
    lines.append(f"- **Action recommandée** : {f['recommended_action']}"
                 + (f" — à solliciter : {f['responsible_party']}" if f.get("responsible_party") else ""))
    if f.get("quality_flags"):
        lines.append(f"- **Alertes qualité** : {', '.join(f['quality_flags'])}")
    if f.get("merged_from"):
        lines.append(f"- **Fusionné avec** : {', '.join(f['merged_from'])} (détecté par {', '.join(f['origin_agents'])})")
    return "\n".join(lines) + "\n"


def table(findings, index):
    out = ["| ID | Titre | Lot | Type | Gravité | Priorité | Statut | Confiance | Preuve |",
           "|---|---|---|---|---|---|---|---|---|"]
    for f in findings:
        out.append(f"| {f['finding_id']} | {md_escape(f['title'])} | {md_escape(f.get('lot'))} | "
                   f"{TYPE_FR.get(f['finding_type'])} | {SEV_FR[f['severity']]} | {f['priority']} | "
                   f"{STATUS_FR[f['status']]} | {CONF_FR[f['confidence']]} | {LEVEL_FR.get(f.get('evidence_level'))} |")
    return "\n".join(out) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--travail", required=True)
    a = ap.parse_args()
    work = Path(a.travail).resolve()
    out = work.parent
    reg = load_json(work / "registre.json")
    if reg is None:
        print_summary({"statut": "echec", "outil": "livrables", "erreur": "registre.json absent : lancer registre.py build"})
        return 0
    man = load_json(work / "manifeste.json", {"documents": []})
    classe = load_json(work / "manifeste_classe.json", {}) or {}
    pre = load_json(work / "preflight.json", {})
    synth = load_json(work / "synthese.json", None)
    etat = load_json(work / "etat_execution.json", {}) or {}
    sel = load_json(work / "selection_verification.json", {}) or {}
    cov = reg.get("couverture", {})
    findings = reg["constats"]
    index = {f["finding_id"]: f for f in findings}
    cls = {d["doc_id"]: d for d in classe.get("documents", [])}
    written = []

    # 1. registre CSV / XLSX
    rows = [row_of(f) for f in findings]
    write_csv(out / "registre_anomalies.csv", CSV_COLS, rows)
    written.append("registre_anomalies.csv")

    # 2. rapprochement économique
    rap_cols = ["type_controle", "controle_id", "statut", "doc_a", "ref_a", "designation_a", "unite_a", "quantite_a",
                "doc_b", "ref_b", "designation_b", "unite_b", "quantite_b", "ecart_abs", "ecart_rel_pct",
                "montant_declare", "montant_recalcule", "ecart_montant", "valorisation_indicative_ht", "commentaire",
                "constats_lies"]
    by_ctrl = defaultdict(list)
    for f in findings:
        c = (f.get("calculation") or {}).get("controle_id")
        if c:
            by_ctrl[c].append(f["finding_id"])
    rap_rows = []
    for p in sorted((work / "controles").glob("rapprochement_*.json")):
        data = load_json(p, {})
        for l in data.get("lignes", []):
            rap_rows.append({"type_controle": "rapprochement_quantites", "controle_id": l["controle_id"],
                             "statut": l.get("statut_appariement") or l["statut"],
                             "doc_a": l.get("doc_a"), "ref_a": l.get("ref_a"), "designation_a": l.get("designation_a"),
                             "unite_a": l.get("unite_a"), "quantite_a": fmt_fr(l.get("quantite_a"), 3),
                             "doc_b": l.get("doc_b"), "ref_b": l.get("ref_b"), "designation_b": l.get("designation_b"),
                             "unite_b": l.get("unite_b"), "quantite_b": fmt_fr(l.get("quantite_b"), 3),
                             "ecart_abs": fmt_fr(l.get("ecart_abs"), 3),
                             "ecart_rel_pct": fmt_fr(l["ecart_rel"] * 100, 2) if l.get("ecart_rel") is not None else "",
                             "valorisation_indicative_ht": fmt_fr(l.get("valorisation_indicative_ht")),
                             "commentaire": l.get("commentaire") or l.get("base_valorisation") or "",
                             "constats_lies": ", ".join(by_ctrl.get(l["controle_id"], []))})
    for p in sorted((work / "controles").glob("dpgf_*.json")):
        data = load_json(p, {})
        for c in data.get("controles", []):
            rap_rows.append({"type_controle": f"controle_{c['type'].lower()}", "controle_id": c["controle_id"],
                             "statut": "anomalie_detectee" if c["type"] in ("ARITH", "TOTAL", "SOUS_TOTAL", "TVA", "TTC")
                             else "a_examiner", "doc_a": data.get("doc_id"), "ref_a": c.get("ref"),
                             "designation_a": c.get("designation"), "unite_a": c.get("unite"),
                             "quantite_a": fmt_fr(c.get("quantite"), 3),
                             "montant_declare": fmt_fr(c.get("montant_declare")),
                             "montant_recalcule": fmt_fr(c.get("montant_recalcule")),
                             "ecart_montant": fmt_fr(c.get("ecart")),
                             "commentaire": f"{c.get('commentaire')} (feuille {c.get('feuille')}, ligne {c.get('ligne')})",
                             "constats_lies": ", ".join(by_ctrl.get(c["controle_id"], []))})
    if not rap_rows:
        rap_rows.append({"type_controle": "aucun", "statut": "non_realise",
                         "commentaire": "aucune pièce quantitative exploitable n'a pu être rapprochée : voir le rapport, "
                                        "section Limites"})
    write_csv(out / "rapprochement_economique.csv", rap_cols, rap_rows)
    written.append("rapprochement_economique.csv")

    # 3. manifeste documentaire
    man_cols = ["doc_id", "chemin_relatif", "format", "nb_pages", "extraction", "type_piece", "intitule", "lots",
                "indice", "date_document", "statut_exploitation", "examine_par", "sha256_court", "remarques"]
    examined = cov.get("documents_examines", {})
    man_rows = []
    for d in man.get("documents", []):
        c = cls.get(d["doc_id"], {})
        man_rows.append({"doc_id": d["doc_id"], "chemin_relatif": d["chemin_relatif"], "format": d["format"],
                         "nb_pages": d.get("nb_pages") or "", "extraction": d.get("extraction"),
                         "type_piece": c.get("type_piece", "non classé"), "intitule": c.get("intitule") or "",
                         "lots": ", ".join(c.get("lots", [])), "indice": c.get("indice") or d.get("indice_detecte_nom") or "",
                         "date_document": c.get("date_document") or "",
                         "statut_exploitation": c.get("statut_exploitation", "non classé"),
                         "examine_par": ", ".join(examined.get(d["doc_id"], [])) or "AUCUN AGENT",
                         "sha256_court": d["sha256"][:12],
                         "remarques": " | ".join(d.get("remarques", []) + ([c["remarque"]] if c.get("remarque") else []))})
    write_csv(out / "manifeste_documentaire.csv", man_cols, man_rows)
    written.append("manifeste_documentaire.csv")

    # 4. questions à arbitrer
    q_by = defaultdict(list)
    for f in findings:
        if f["status"] == "refute":
            continue
        for q in f.get("open_questions", []) or []:
            q_by[role(f.get("responsible_party"))].append((q, f))
        arb = f.get("arbitrage") or {}
        for q in arb.get("questions_pour_moe", []) or []:
            q_by["Architecte / maîtrise d'œuvre"].append((q, f))
        if f["status"] in ("non_resolu", "conteste") and not f.get("open_questions") and not arb.get("questions_pour_moe"):
            doutes = [v["justification"] for v in f.get("verifications", []) if v["verdict"] in ("insuffisant", "refute")]
            q = f"Confirmer ou infirmer : {f['claim']}" + (f" (point non établi : {doutes[0]})" if doutes else "")
            q_by[role(f.get("responsible_party"))].append((q, f))
    ql = ["# Questions et arbitrages à obtenir", "",
          "Chaque question est rattachée au constat qui la motive. Une réponse est nécessaire pour lever "
          "l'incertitude indiquée ; tant qu'elle manque, le constat reste au statut affiché.", ""]
    n = 0
    for who in sorted(q_by):
        ql.append(f"## Destinataire : {who}")
        ql.append("")
        items = sorted(q_by[who], key=lambda qf: (SEV_ORDER.get(qf[1]["severity"], 9), qf[1]["priority"]))
        for q, f in items:
            n += 1
            srcs = "; ".join(src_str(s) for s in f.get("sources", [])[:2]) or "sans source"
            ql += [f"### Q{n:03d} — {q}",
                   f"- **Constat** : {f['finding_id']} — {f['title']} ({TYPE_FR[f['finding_type']]}, "
                   f"{SEV_FR[f['severity']]}, {f['priority']}, statut {STATUS_FR[f['status']].lower()})",
                   f"- **Contexte** : {f['claim']}",
                   f"- **Acteur cité par l'agent** : {f.get('responsible_party') or 'non précisé'}",
                   f"- **Pièces concernées** : {srcs}",
                   f"- **Pourquoi une réponse est nécessaire** : {f.get('severity_rationale')}"
                   + (f" — échéance : {f['decision_deadline']}" if f.get("decision_deadline") else ""), ""]
    if classe.get("documents_manquants"):
        ql += ["## Pièces à transmettre (citées mais absentes du dossier)", ""]
        for m in classe["documents_manquants"]:
            n += 1
            txt = m if isinstance(m, str) else json.dumps(m, ensure_ascii=False)
            ql.append(f"- Q{n:03d} — {txt}")
        ql.append("")
    if classe.get("questions_cadrage"):
        ql += ["## Questions de cadrage (documentaliste)", ""]
        for q in classe["questions_cadrage"]:
            n += 1
            ql.append(f"- Q{n:03d} — {q if isinstance(q, str) else json.dumps(q, ensure_ascii=False)}")
        ql.append("")
    ecartees = []
    for p in sorted((work / "echanges").glob("r*/a_traiter.json")):
        for e in load_json(p, {}).get("ecartees", []):
            if e.get("question"):
                ecartees.append(e)
    if ecartees:
        ql += ["## Questions internes non traitées par le workflow", "",
               "Ces questions posées entre agents n'ont pas reçu de réponse (agent inactif ou plafond atteint). "
               "Elles restent ouvertes.", ""]
        for e in ecartees:
            n += 1
            ql.append(f"- Q{n:03d} — {e['question']} (destinataire prévu : {e.get('destinataire')}, motif : {e['motif']})")
        ql.append("")
    if n == 0:
        ql.append("Aucune question n'a été formulée.")
    (out / "questions_a_arbitrer.md").write_text("\n".join(ql) + "\n", encoding="utf-8")
    written.append("questions_a_arbitrer.md")

    # 5. risques et interfaces
    inter = [f for f in findings if f["status"] != "refute" and (
        "interfaces" in (f.get("severity_criteria") or []) or f.get("agent_origine") in ("synthese_interlots", "opc")
        or "interface" in (f.get("discipline") or "").lower() or "blocage_chantier" in (f.get("severity_criteria") or []))]
    rl = ["# Risques de chantier, interfaces et dépendances", "",
          "Les contraintes ci-dessous sont issues des pièces analysées. Les durées et délais ne sont jamais "
          "inventés : quand une échéance figure, elle provient d'une pièce ou de la nature de la décision.", ""]
    if inter:
        rl.append(table(sorted(inter, key=lambda f: (SEV_ORDER[f["severity"]], f["priority"])), index))
        for f in sorted(inter, key=lambda f: (SEV_ORDER[f["severity"]], f["priority"])):
            rl.append(finding_block(f, index))
            if f.get("dependencies"):
                rl.append(f"Dépend de : {', '.join(fid_label(d, index) for d in f['dependencies'])}\n")
    else:
        rl.append("Aucun risque d'interface ou de séquencement n'a été établi à partir des pièces disponibles.")
    deadlines = [f for f in findings if f["status"] != "refute" and f.get("decision_deadline")]
    if deadlines:
        rl += ["", "## Décisions à prendre et échéances", "", "| Constat | Décision attendue | Échéance | Statut |",
               "|---|---|---|---|"]
        for f in sorted(deadlines, key=lambda f: (f["priority"], SEV_ORDER[f["severity"]])):
            rl.append(f"| {f['finding_id']} | {md_escape(f['recommended_action'])} | {md_escape(f['decision_deadline'])} "
                      f"| {STATUS_FR[f['status']]} |")
    (out / "risques_et_interfaces.md").write_text("\n".join(rl) + "\n", encoding="utf-8")
    written.append("risques_et_interfaces.md")

    # 6. journal des contrôles
    jl = ["# Journal des contrôles et de l'exécution", ""]
    jl += ["## Outils déterministes exécutés", "", "| Horodatage | Outil | Statut | Détails |", "|---|---|---|---|"]
    jpath = work / "journal_controles.jsonl"
    if jpath.exists():
        for line in jpath.read_text(encoding="utf-8").splitlines():
            try:
                r = json.loads(line)
            except json.JSONDecodeError:
                continue
            jl.append(f"| {r['horodatage']} | {r['outil']} | {r['statut']} | "
                      f"{md_escape(json.dumps(r.get('details', {}), ensure_ascii=False)[:300])} |")
    jl += ["", "## Agents lancés par le workflow", ""]
    if etat.get("agents"):
        jl += ["| Agent | Phase | Résultat | Remarque |", "|---|---|---|---|"]
        for ag in etat["agents"]:
            jl.append(f"| {ag.get('label')} | {ag.get('phase')} | {ag.get('statut')} | {md_escape(ag.get('remarque'))} |")
    else:
        jl.append("État d'exécution non transmis par l'orchestrateur.")
    for key, title in (("plan", "Plan d'analyse retenu"), ("limites", "Limites et réductions de périmètre"),
                       ("divergences", "Divergences détectées entre résultats d'agents et fichiers")):
        if etat.get(key):
            jl += ["", f"## {title}", ""]
            val = etat[key]
            if isinstance(val, list):
                jl += [f"- {x if isinstance(x, str) else json.dumps(x, ensure_ascii=False)}" for x in val]
            else:
                jl.append("```json\n" + json.dumps(val, ensure_ascii=False, indent=2) + "\n```")
    if reg.get("fichiers_invalides"):
        jl += ["", "## Fichiers d'agents rejetés ou partiellement rejetés (non conformes au contrat de données)", ""]
        for inv in reg["fichiers_invalides"]:
            jl.append(f"- {inv['fichier']} : {md_escape(json.dumps(inv.get('erreurs') or inv.get('constats_rejetes'), ensure_ascii=False)[:400])}")
    (out / "journal_controles.md").write_text("\n".join(jl) + "\n", encoding="utf-8")
    written.append("journal_controles.md")

    # 7. rapport exécutif
    st = Counter(f["status"] for f in findings)
    etablis = [f for f in findings if f["status"] == "corrobore"]
    averif = [f for f in findings if f["status"] == "a_verifier"]
    disputes = [f for f in findings if f["status"] in ("conteste", "non_resolu")]
    refutes = [f for f in findings if f["status"] == "refute"]
    src = pre.get("source", {}).get("chemin", man.get("source"))
    R = [f"# Rapport d'audit du dossier de construction", "",
         f"- **Dossier analysé** : `{src}`",
         f"- **Date d'analyse** : {pre.get('horodatage', 'non disponible')}",
         f"- **Mode** : {etat.get('mode', pre.get('mode', 'non précisé'))}",
         f"- **Documents inventoriés** : {len(man.get('documents', []))} — constats au registre : {len(findings)} "
         f"(corroborés {st.get('corrobore', 0)}, à vérifier {st.get('a_verifier', 0)}, contestés {st.get('conteste', 0)}, "
         f"non résolus {st.get('non_resolu', 0)}, réfutés {st.get('refute', 0)})", "",
         "> Ce rapport est une aide à l'analyse produite par un système multi-agents. Il ne remplace ni l'avis "
         "d'un professionnel qualifié, ni une note de calcul, ni un contrôle réglementaire. Chaque constat porte "
         "son statut de vérification : seuls les constats « corroborés » ont résisté à une vérification indépendante.", ""]
    R += ["## 1. Synthèse exécutive", ""]
    if synth:
        R += [synth.get("appreciation_globale", ""), "", "**Messages clés**", ""]
        for m in synth.get("messages_cles", []):
            R.append(f"- {m['texte']} — constats : {', '.join(fid_label(x, index) for x in m.get('constats', []))}")
        R += ["", "**Décisions prioritaires à prendre**", ""]
        for d in synth.get("decisions_prioritaires", []):
            R.append(f"- {d['decision']}" + (f" — acteur : {d['acteur']}" if d.get("acteur") else "")
                     + (f" — échéance : {d['echeance']}" if d.get("echeance") else "")
                     + f" — constats : {', '.join(fid_label(x, index) for x in d.get('constats', []))}")
        for w in synth.get("avertissements", []) or []:
            R.append(f"- ⚠ {w}")
    else:
        R.append("Synthèse rédigée non disponible (agent rédacteur absent ou en échec) : voir les sections suivantes, "
                 "générées directement depuis le registre.")
    R += ["", "## 2. Périmètre et documents analysés", ""]
    fmt_counts = Counter(d["format"] for d in man.get("documents", []))
    R.append("Formats : " + ", ".join(f"{k} ({v})" for k, v in sorted(fmt_counts.items())))
    R += ["", "| Doc | Fichier | Type de pièce | Lots | Indice | Exploitation | Examiné par |", "|---|---|---|---|---|---|---|"]
    for r in man_rows:
        R.append(f"| {r['doc_id']} | {md_escape(r['chemin_relatif'])} | {md_escape(r['type_piece'])} | {md_escape(r['lots'])} | "
                 f"{md_escape(r['indice'])} | {md_escape(r['statut_exploitation'])} / {r['extraction']} | {md_escape(r['examine_par'])} |")
    R += ["", "## 3. Limites de l'analyse", "", "<!-- section:limites -->"]
    lim = []
    for w in pre.get("avertissements", []):
        lim.append(w)
    for d in cov.get("documents_jamais_examines", []):
        lim.append(f"{d['doc_id']} ({d['nom']}) n'a été examiné par aucun agent")
    for d in cov.get("documents_non_exploitables", []):
        lim.append(f"{d['doc_id']} ({d['nom']}) non exploitable par les outils : {d['extraction']}")
    for x in cov.get("agents_non_termines", []):
        lim.append(f"analyse {x['statut_analyse']} pour l'agent {x['agent']} ({x.get('phase')}) : {x.get('motif') or 'motif non précisé'}")
    for ag in etat.get("agents", []):
        if ag.get("statut") in ("echec", "vide", "abandon"):
            lim.append(f"agent {ag['label']} ({ag.get('phase')}) : {ag['statut']} — {ag.get('remarque') or ''}")
    for x in etat.get("limites", []) or []:
        lim.append(x if isinstance(x, str) else json.dumps(x, ensure_ascii=False))
    if sel.get("non_verifies_plafond"):
        lim.append(f"plafond de vérification atteint : {len(sel['non_verifies_plafond'])} constat(s) non vérifiés "
                   f"({', '.join(sel['non_verifies_plafond'][:15])})")
    if cov.get("constats_graves_non_verifies"):
        lim.append("constats graves restés sans vérification indépendante : " + ", ".join(cov["constats_graves_non_verifies"]))
    if cov.get("lots_sans_analyse_explicite"):
        lim.append("lots mentionnés dans les pièces sans analyse explicite : " + ", ".join(
            f"lot {l['lot']} ({', '.join(l['intitules'][:2])})" for l in cov["lots_sans_analyse_explicite"]))
    for a_name, a_val in reg.get("agents", {}).items():
        for x in a_val.get("limites", []):
            lim.append(f"[{a_name}] {x}")
    R += [f"- {x}" for x in lim] or ["- Aucune limite particulière déclarée."]
    R += ["<!-- /section:limites -->", "", "## 4. Constats établis (corroborés)", "", "<!-- section:etablis -->"]
    if etablis:
        R.append(table(etablis, index))
        for f in etablis:
            R.append(finding_block(f, index))
    else:
        R.append("Aucun constat n'a été corroboré par une vérification indépendante.")
    R += ["<!-- /section:etablis -->", "", "## 5. Constats non vérifiés de façon indépendante (à vérifier)", ""]
    if averif:
        R.append("Ces constats reposent sur l'analyse d'un seul agent ; ils n'ont été ni confirmés ni réfutés.")
        R.append(table(averif, index))
        for f in averif:
            if f["severity"] in ("critique", "majeure"):
                R.append(finding_block(f, index))
    else:
        R.append("Aucun.")
    R += ["", "## 6. Désaccords, constats contestés et non résolus", "", "<!-- section:desaccords -->"]
    if disputes:
        for f in disputes:
            R.append(finding_block(f, index))
            for e in f.get("echanges", []):
                R.append(f"  - Échange tour {e['round']} — question de {e.get('de')}, réponse de {e['a']}"
                         f"{'' if e.get('independant') else ' (co-auteur, non indépendant)'} : **{e['verdict']}** — {e['justification']}")
            for v in f.get("verifications", []):
                R.append(f"  - Vérification ({v['lentille']}) : **{v['verdict']}** — {v['justification']}")
            arb = f.get("arbitrage")
            if arb:
                R.append(f"  - Arbitrage : **{arb['decision']}** — {arb['motif']}")
                for pos in arb.get("positions", []):
                    R.append(f"    - Position {pos['acteur']} : {pos['position']}" + (f" ({pos['argument']})" if pos.get("argument") else ""))
            R.append("")
    else:
        R.append("Aucun désaccord persistant.")
    R += ["<!-- /section:desaccords -->", "", "## 7. Risques économiques", ""]
    eco = [f for f in findings if f["status"] != "refute" and (f.get("economic_impact") or {}).get("quantifie")]
    if eco:
        R += ["Impacts chiffrés (non additionnables sans examen : certains constats peuvent se recouper).", "",
              "| Constat | Montant HT | Base de calcul | Statut |", "|---|---|---|---|"]
        for f in sorted(eco, key=lambda f: -abs(f["economic_impact"].get("montant_ht") or 0)):
            R.append(f"| {f['finding_id']} {md_escape(f['title'])} | {eur(f['economic_impact'].get('montant_ht'))} | "
                     f"{md_escape(f['economic_impact'].get('base'))} | {STATUS_FR[f['status']]} |")
    unq = [f for f in findings if f["status"] != "refute" and not (f.get("economic_impact") or {}).get("quantifie")
           and "economique" in (f.get("severity_criteria") or [])]
    if unq:
        R += ["", "Impacts économiques potentiels non quantifiables avec les données disponibles : "
              + ", ".join(f"{f['finding_id']}" for f in unq)]
    if not eco and not unq:
        R.append("Aucun risque économique établi.")
    R += ["", "## 8. Incohérences et interfaces entre lots", "",
          f"Voir le détail dans `risques_et_interfaces.md` ({len(inter)} constat(s))."]
    if inter:
        R.append(table(inter, index))
    R += ["", "## 9. Constats réfutés (transparence)", ""]
    if refutes:
        R += ["| ID | Titre | Motif du rejet |", "|---|---|---|"]
        for f in refutes:
            R.append(f"| {f['finding_id']} | {md_escape(f['title'])} | {md_escape(f.get('motif_statut'))} |")
    else:
        R.append("Aucun.")
    R += ["", "## 10. Recommandations d'action", ""]
    for f in sorted([f for f in findings if f["status"] in ("corrobore", "a_verifier", "conteste", "non_resolu")],
                    key=lambda f: (f["priority"], SEV_ORDER[f["severity"]]))[:30]:
        R.append(f"- **{f['priority']}** — {f['recommended_action']} ({fid_label(f['finding_id'], index)})")
    R += ["", "## 11. Méthode et lecture", "",
          "- **Statuts** : corroboré (confirmé par un vérificateur, un pair indépendant ou un contrôle déterministe, "
          "sans réfutation) ; à vérifier (un seul agent) ; contesté (confirmations et réfutations) ; non résolu "
          "(désaccord persistant ou preuve jugée insuffisante) ; réfuté.",
          "- **Gravité ≠ confiance** : la gravité mesure les conséquences si le constat est vrai ; la confiance mesure "
          "la solidité de la preuve. Un risque grave peu certain appelle une vérification rapide, pas un abandon.",
          "- **Niveau de preuve** calculé automatiquement : contrôle déterministe, citation retrouvée mot pour mot "
          "dans la pièce, observation localisée, ou non étayé.",
          "- Fichiers : `registre_anomalies.csv`, `rapprochement_economique.csv`, `questions_a_arbitrer.md`, "
          "`risques_et_interfaces.md`, `manifeste_documentaire.csv`, `journal_controles.md`, dossier `_travail/` "
          "(traces complètes : constats par agent, échanges, vérifications, arbitrages)."]
    (out / "rapport_audit_dossier_btp.md").write_text("\n".join(R) + "\n", encoding="utf-8")
    written.append("rapport_audit_dossier_btp.md")

    # XLSX optionnel
    try:
        import openpyxl  # type: ignore
        from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Registre"
        ws.append(CSV_COLS)
        for r in rows:
            ws.append([r.get(c) for c in CSV_COLS])
        ws2 = wb.create_sheet("Rapprochement")
        ws2.append(rap_cols)
        for r in rap_rows:
            ws2.append([r.get(c) for c in rap_cols])
        ws3 = wb.create_sheet("Manifeste")
        ws3.append(man_cols)
        for r in man_rows:
            ws3.append([r.get(c) for c in man_cols])
        for w in (ws, ws2, ws3):
            for c in w[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
                c.alignment = Alignment(wrap_text=True, vertical="top")
            w.freeze_panes = "B2"
            w.auto_filter.ref = w.dimensions
            for col in w.columns:
                w.column_dimensions[col[0].column_letter].width = 22
        wb.save(out / "registre_anomalies.xlsx")
        written.append("registre_anomalies.xlsx")
    except ImportError:
        pass

    journal(work, "livrables", "ok", {"fichiers": written})
    print_summary({"statut": "ok", "outil": "livrables", "sortie": str(out), "fichiers": written,
                   "nb_constats": len(findings), "par_statut": dict(st)})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
